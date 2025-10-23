"""
Quote Builder for constructing quote documents.

This module implements the Builder pattern for creating
complex quote documents step by step.
"""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from constants import (
    DEFAULT_FOOTER_TEXT,
    SUCCESS_MESSAGE_TEMPLATE,
    EDITABLE_MESSAGE,
    FONT_SIZE_SMALL
)
from sections import (
    add_header_section,
    add_customer_info_section,
    add_products_table,
    add_summary_section,
    add_signature_section,
    add_banking_info,
    add_terms_and_conditions
)


class QuoteBuilder:
    """
    Builder pattern implementation for constructing quote documents.

    This builder provides a fluent interface for creating complex
    quote documents step by step. Each method returns self to allow
    method chaining.

    Example usage:
        builder = QuoteBuilder()
        doc = (builder.with_header()
                      .with_customer_info(customer_data, quote_data)
                      .with_products(products)
                      .with_summary(summary)
                      .with_signature()
                      .with_banking_info()
                      .with_terms()
                      .with_footer()
                      .build())
    """

    def __init__(self):
        """Initialize a new QuoteBuilder with a blank document."""
        self.doc = Document()

    def with_header(self):
        """
        Add company header information including logo and
        contact details.

        Returns:
            self: Returns the builder instance for method chaining
        """
        add_header_section(self.doc)
        return self

    def with_customer_info(self, customer_data, quote_data):
        """
        Add customer information and quote details.

        Args:
            customer_data: Dictionary containing customer information
            quote_data: Dictionary containing quote metadata

        Returns:
            self: Returns the builder instance for method chaining
        """
        add_customer_info_section(
            self.doc, customer_data, quote_data
        )
        return self

    def with_products(self, products):
        """
        Add products table with pricing details.

        Args:
            products: List of product dictionaries

        Returns:
            self: Returns the builder instance for method chaining
        """
        add_products_table(self.doc, products)
        self.doc.add_paragraph()  # Spacer
        return self

    def with_summary(self, summary):
        """
        Add financial summary section.

        Args:
            summary: Dictionary containing subtotal, discount, total,
                     anticipo

        Returns:
            self: Returns the builder instance for method chaining
        """
        add_summary_section(self.doc, summary)
        return self

    def with_signature(self):
        """
        Add signature section for customer acceptance.

        Returns:
            self: Returns the builder instance for method chaining
        """
        add_signature_section(self.doc)
        return self

    def with_banking_info(self):
        """
        Add banking information table.

        Returns:
            self: Returns the builder instance for method chaining
        """
        add_banking_info(self.doc)
        return self

    def with_terms(self):
        """
        Add terms and conditions section.

        Returns:
            self: Returns the builder instance for method chaining
        """
        add_terms_and_conditions(self.doc)
        return self

    def with_footer(self, footer_text=None):
        """
        Add footer with location information.

        Args:
            footer_text: Optional custom footer text. If None,
                         uses default locations.

        Returns:
            self: Returns the builder instance for method chaining
        """
        if footer_text is None:
            footer_text = DEFAULT_FOOTER_TEXT

        section = self.doc.sections[0]
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(footer_text)
        run.font.size = Pt(FONT_SIZE_SMALL)
        return self

    def build(self):
        """
        Build and return the final document.

        Returns:
            Document: The constructed python-docx Document object
        """
        return self.doc

    def save(self, output_filename):
        """
        Build and save the document to a file.

        Args:
            output_filename: Name of the output file

        Returns:
            str: The output filename
        """
        self.doc.save(output_filename)
        print(SUCCESS_MESSAGE_TEMPLATE.format(output_filename))
        print(EDITABLE_MESSAGE)
        return output_filename
