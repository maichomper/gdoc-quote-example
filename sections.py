"""
Document section builders for quote generation.

This module contains functions to build different sections
of the quote document (header, customer info, products, etc.).
"""

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from constants import (
    COMPANY_WEBSITE,
    COMPANY_NAME,
    COMPANY_LOCATION,
    COMPANY_ADDRESS,
    COMPANY_PHONE,
    COMPANY_EMAIL,
    LABEL_NOMBRE,
    LABEL_TELEFONO,
    LABEL_DIRECCION,
    LABEL_MAIL,
    LABEL_NO_COTIZACION,
    LABEL_FECHA,
    LABEL_INFO_CONTACTO,
    PRODUCT_TABLE_HEADERS,
    SUMMARY_SUBTOTAL,
    SUMMARY_DESCUENTO,
    SUMMARY_TOTAL,
    SUMMARY_ANTICIPO_TEMPLATE,
    SIGNATURE_LINE,
    SIGNATURE_TEXT_NAME,
    SIGNATURE_TEXT_ACCEPT,
    BANK_BENEFICIARY,
    BANKING_INFO,
    BANKING_HEADER,
    BANKING_ROW_BENEFICIARIO,
    BANKING_ROW_CUENTA,
    BANKING_ROW_CLABE,
    TERMS_HEADING,
    TERMS_SECTION_PAGOS,
    TERMS_PAGOS_LIST,
    TERMS_ADDITIONAL_NOTE,
    CELL_BG_BLACK,
    FONT_SIZE_SMALL,
    FONT_SIZE_NORMAL,
    FONT_SIZE_MEDIUM,
    FONT_SIZE_LARGE,
    FONT_SIZE_XLARGE,
    FONT_SIZE_HEADING,
)
from document_styles import set_cell_background
from image_utils import generate_placeholder_image


def add_header_section(doc):
    """Add company header information."""
    # Add logo
    logo_para = doc.add_paragraph()
    logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        logo_image = generate_placeholder_image(
            400, 200, seed=1, text="LOGO"
        )
        logo_run = logo_para.add_run()
        logo_run.add_picture(logo_image, width=Inches(2.5))
    except Exception as e:
        print(f"Warning: Could not add logo image: {e}")
        logo_para.add_run('[LOGO]')

    # Company name and website
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = header.add_run(f'{COMPANY_WEBSITE}\n')
    run.font.size = Pt(FONT_SIZE_MEDIUM)

    run = header.add_run(f'{COMPANY_NAME}\n\n')
    run.font.size = Pt(FONT_SIZE_LARGE)
    run.font.bold = True

    run = header.add_run(f'{COMPANY_LOCATION}\n')
    run.font.size = Pt(FONT_SIZE_XLARGE)
    run.font.bold = True

    run = header.add_run(f'{COMPANY_ADDRESS}\n')
    run.font.size = Pt(FONT_SIZE_NORMAL)

    run = header.add_run(f'{COMPANY_PHONE}\n')
    run.font.size = Pt(FONT_SIZE_NORMAL)

    run = header.add_run(COMPANY_EMAIL)
    run.font.size = Pt(FONT_SIZE_NORMAL)

    doc.add_paragraph()  # Spacer


def add_customer_info_section(doc, customer_data, quote_data):
    """Add customer and quote information in two columns."""
    table = doc.add_table(rows=5, cols=2)
    table.autofit = False
    table.allow_autofit = False

    # Left column - Customer info
    cells = table.rows[0].cells
    p = cells[0].paragraphs[0]
    run = p.add_run(LABEL_NOMBRE)
    run.font.bold = True
    run = p.add_run(customer_data['nombre'])

    # Right column - Quote number
    p = cells[1].paragraphs[0]
    run = p.add_run(LABEL_NO_COTIZACION)
    run.font.bold = True
    run = p.add_run(quote_data['numero'])

    # Row 2
    cells = table.rows[1].cells
    p = cells[0].paragraphs[0]
    run = p.add_run(LABEL_TELEFONO)
    run.font.bold = True
    run = p.add_run(customer_data['telefono'])

    p = cells[1].paragraphs[0]
    run = p.add_run(LABEL_FECHA)
    run.font.bold = True
    run = p.add_run(quote_data['fecha'])

    # Row 3
    cells = table.rows[2].cells
    p = cells[0].paragraphs[0]
    run = p.add_run(LABEL_DIRECCION)
    run.font.bold = True
    run = p.add_run(customer_data.get('direccion', ''))

    p = cells[1].paragraphs[0]
    run = p.add_run(LABEL_INFO_CONTACTO)
    run.font.bold = True

    # Row 4
    cells = table.rows[3].cells
    p = cells[0].paragraphs[0]
    run = p.add_run(LABEL_MAIL)
    run.font.bold = True
    run = p.add_run(customer_data.get('mail', ''))

    p = cells[1].paragraphs[0]
    p.add_run(quote_data['contacto'])

    # Row 5
    cells = table.rows[4].cells
    cells[0].text = ''
    p = cells[1].paragraphs[0]
    p.add_run(quote_data.get('email_contacto', ''))

    doc.add_paragraph()  # Spacer


def add_products_table(doc, products):
    """Add products table with all details."""
    # Create table with header row
    table = doc.add_table(rows=1, cols=9)
    table.style = 'Light Grid Accent 1'

    # Header row
    header_cells = table.rows[0].cells

    for i, header in enumerate(PRODUCT_TABLE_HEADERS):
        cell = header_cells[i]
        set_cell_background(cell, CELL_BG_BLACK)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.font.bold = True
        run.font.size = Pt(FONT_SIZE_NORMAL)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Add product rows
    for idx, product in enumerate(products):
        row_cells = table.add_row().cells

        # Marca
        row_cells[0].text = product['marca']

        # Descripción
        row_cells[1].text = product['descripcion']

        # Imagen Referencia
        p = row_cells[2].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            # Use different seed for each product
            product_image = generate_placeholder_image(
                300, 300, seed=idx + 10, text=f"P{idx+1}"
            )
            img_run = p.add_run()
            img_run.add_picture(product_image, width=Inches(1.2))
        except Exception as e:
            print(f"Warning: Could not add product image {idx}: {e}")
            p.add_run('[IMAGEN]')

        # Código
        row_cells[3].text = product['codigo']

        # Piezas
        p = row_cells[4].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(str(product['piezas']))

        # Precio Unitario
        p = row_cells[5].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run(f"$ {product['precio_unitario']:,.2f}")

        # Descuento
        p = row_cells[6].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        discount_pct = product['descuento_pct']
        discount_amt = product['descuento_monto']
        p.add_run(f"{discount_pct:.1f}%\n")
        p.add_run(f"$ {discount_amt:,.2f}")

        # Precio con Descuento
        p = row_cells[7].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run(f"$ {product['precio_descuento']:,.2f}")

        # Total
        p = row_cells[8].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run(f"$ {product['total']:,.2f}")

    return table


def add_summary_section(doc, summary):
    """Add financial summary table."""
    # Create a right-aligned table for the summary
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Set column widths
    table.columns[0].width = Inches(4.5)
    table.columns[1].width = Inches(2.0)

    anticipo_label = SUMMARY_ANTICIPO_TEMPLATE.format(
        summary['anticipo_pct']
    )

    summary_items = [
        (SUMMARY_SUBTOTAL, summary['subtotal']),
        (SUMMARY_DESCUENTO, summary['descuento']),
        (SUMMARY_TOTAL, summary['total']),
        (anticipo_label, summary['anticipo'])
    ]

    for i, (label, value) in enumerate(summary_items):
        row = table.rows[i]

        # Label cell
        set_cell_background(row.cells[0], CELL_BG_BLACK)
        p = row.cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(label)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

        # Value cell
        p = row.cells[1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(f"$ {value:,.2f}")
        run.font.bold = True


def add_signature_section(doc):
    """Add signature section."""
    doc.add_paragraph()
    doc.add_paragraph(SIGNATURE_LINE)

    p = doc.add_paragraph(SIGNATURE_TEXT_NAME)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.bold = True

    p = doc.add_paragraph(SIGNATURE_TEXT_ACCEPT)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.bold = True


def add_banking_info(doc):
    """Add banking information table."""
    doc.add_page_break()

    table = doc.add_table(rows=4, cols=4)
    table.style = 'Light Grid Accent 1'

    # Header row
    headers = [BANKING_HEADER] + BANKING_INFO['banks']
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        if i == 0:
            set_cell_background(cell, CELL_BG_BLACK)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.font.bold = True
        if i == 0:
            run.font.color.rgb = RGBColor(255, 255, 255)

    # Beneficiario
    _add_banking_row(
        table, 1, BANKING_ROW_BENEFICIARIO,
        [BANK_BENEFICIARY] * 3
    )

    # Número de cuenta
    _add_banking_row(
        table, 2, BANKING_ROW_CUENTA,
        BANKING_INFO['accounts']
    )

    # CLABE
    _add_banking_row(
        table, 3, BANKING_ROW_CLABE,
        BANKING_INFO['clabe']
    )


def _add_banking_row(table, row_idx, label, values):
    """
    Helper function to add a row to banking info table.

    Args:
        table: The table object
        row_idx: Row index
        label: Label for first cell
        values: List of values for remaining cells
    """
    row = table.rows[row_idx]
    row.cells[0].text = label
    set_cell_background(row.cells[0], CELL_BG_BLACK)
    row.cells[0].paragraphs[0].runs[0].font.color.rgb = (
        RGBColor(255, 255, 255)
    )
    row.cells[0].paragraphs[0].runs[0].font.bold = True

    for i, value in enumerate(values, start=1):
        row.cells[i].text = value


def add_terms_and_conditions(doc):
    """Add commercial policies and terms."""
    doc.add_paragraph()

    heading = doc.add_paragraph(TERMS_HEADING)
    heading.runs[0].font.bold = True
    heading.runs[0].font.size = Pt(FONT_SIZE_HEADING)

    # PAGOS section
    section = doc.add_paragraph(TERMS_SECTION_PAGOS)
    section.runs[0].font.bold = True
    section.runs[0].font.size = Pt(FONT_SIZE_LARGE)

    for term in TERMS_PAGOS_LIST:
        p = doc.add_paragraph(term, style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.25)

    # Note: Additional terms sections would be added similarly
    doc.add_paragraph()
    p = doc.add_paragraph(TERMS_ADDITIONAL_NOTE)
    p.runs[0].font.italic = True
